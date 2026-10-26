"""
Reading Word (.docx) and PowerPoint (.pptx) files — text AND pictures.

    extract_any(path)        -> [{"source", "page", "text"}]   like step1's extract_file()
    office_images(path, n)   -> [{"page", "png", "context", "caption", ...}]
                                same shape as images.extract_images(), so the
                                vision pipeline (describe, store, index, restore)
                                works for Office files exactly as for PDFs.

WORD (.docx)
    Paragraphs and tables, in document order. Headings are kept as their own
    lines so chunks start at sensible places. A .docx has no fixed pages, but
    Word records where pages broke when the file was last saved; those marks
    give approximate page numbers for citations. (Files made by other tools
    may not have them: then citations name only the file.)

POWERPOINT (.pptx)
    One "page" per slide (shown as "slide N" in citations): the title, every
    text box, grouped shapes, tables, and the speaker notes.

PICTURES
    Embedded pictures (photos, screenshots, diagrams saved as images). Tiny
    images and ones repeated on many slides (logos) are skipped. Vector
    formats Word/PowerPoint use for some clip-art (EMF/WMF) can't be read and
    are skipped. Charts made with PowerPoint's own chart tool are not pictures:
    their title and data labels are usually not extracted either.
"""
import hashlib
import re
from collections import Counter
from pathlib import Path

OFFICE = {".docx", ".pptx"}
MIN_IMAGE_PX = 120            # narrower/shorter pictures are icons and bullets
REPEATED_ON = 3               # a picture on this many slides/pages is a logo or background

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def _clean(text: str) -> str:
    text = re.sub(r"[ \t ]+", " ", text or "")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_any(path: Path) -> list[dict]:
    """Text of any supported file, as a list of pages."""
    from vector_store import extract_file

    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _docx_pages(path)
    if suffix == ".pptx":
        return _pptx_pages(path)
    return extract_file(path)


# =========================================================================== Word
def _docx_blocks(doc):
    """Paragraphs and tables in reading order: ("p", paragraph) / ("t", table)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if child.tag == f"{_W}p":
            yield "p", Paragraph(child, doc)
        elif child.tag == f"{_W}tbl":
            yield "t", Table(child, doc)


def _page_breaks(element) -> int:
    """Page breaks inside one paragraph: Word's 'page ended here' marks + manual breaks."""
    rendered = len(element.findall(f".//{_W}lastRenderedPageBreak"))
    manual = sum(1 for br in element.findall(f".//{_W}br") if br.get(f"{_W}type") == "page")
    return max(rendered, manual)


def _table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = list(dict.fromkeys(c.text.strip() for c in row.cells))     # merged cells repeat
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _docx_pages(path: Path) -> list[dict]:
    import docx

    doc = docx.Document(str(path))
    pages: dict[int, list[str]] = {1: []}
    page, has_breaks = 1, False
    for kind, block in _docx_blocks(doc):
        breaks = _page_breaks(block._element)
        if breaks:
            has_breaks = True
            page += breaks
            pages.setdefault(page, [])
        if kind == "p":
            text = block.text.strip()
            if text:
                style = (block.style.name or "").lower() if block.style is not None else ""
                pages[page].append(f"\n{text}\n" if style.startswith(("heading", "title")) else text)
        else:
            text = _table_text(block)
            if text:
                pages[page].append(f"\n{text}\n")
    if not has_breaks:
        text = _clean("\n".join(t for p in pages.values() for t in p))
        return [{"source": path.name, "page": None, "text": text}] if text else []
    return [{"source": path.name, "page": n, "text": _clean("\n".join(lines))}
            for n, lines in sorted(pages.items()) if "".join(lines).strip()]


# =========================================================================== PowerPoint
def _shape_texts(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_texts(shape.shapes)
        elif getattr(shape, "has_table", False) and shape.has_table:
            rows = [" | ".join(c.text.strip() for c in r.cells) for r in shape.table.rows]
            yield "\n".join(r for r in rows if r.strip(" |"))
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                yield text


def _pptx_pages(path: Path) -> list[dict]:
    from pptx import Presentation

    out = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        parts = list(_shape_texts(slide.shapes))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append(f"Speaker notes: {notes}")
        text = _clean("\n".join(parts))
        if text:
            out.append({"source": path.name, "page": n, "text": text})
    return out


# =========================================================================== pictures
def _pptx_pictures(path: Path):
    """(slide number, image bytes, slide text) for every picture on every slide."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def pictures(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from pictures(shape.shapes)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                try:
                    yield shape.image.blob
                except Exception:
                    continue

    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        context = " ".join(_clean(" ".join(_shape_texts(slide.shapes))).split())[:300]
        for blob in pictures(slide.shapes):
            yield n, blob, context, ""


def _docx_pictures(path: Path):
    """(approximate page, image bytes, nearby text, caption) for every picture in the document."""
    import docx

    doc = docx.Document(str(path))
    blocks = [(k, b) for k, b in _docx_blocks(doc)]
    page = 1
    for i, (kind, block) in enumerate(blocks):
        page += _page_breaks(block._element)
        for blip in block._element.iter(_A_BLIP):
            rid = blip.get(_R_EMBED)
            part = doc.part.related_parts.get(rid) if rid else None
            if part is None:
                continue
            near = [b.text.strip() for k, b in blocks[max(0, i - 2): i + 3] if k == "p" and b.text.strip()]
            caption = next((t for t in near if re.match(r"(fig(ure)?\.?|image|photo|chart|diagram)\s*\d*", t, re.I)
                            and len(t) < 300), "")
            yield page, part.blob, " ".join(" ".join(near).split())[:300], caption


def office_images(path: Path, max_images: int) -> list[dict]:
    """Pictures from a .docx/.pptx in the same format images.extract_images() returns."""
    from PIL import Image

    from images import _to_png

    path = Path(path)
    source = _pptx_pictures(path) if path.suffix.lower() == ".pptx" else _docx_pictures(path)
    items = list(source)
    counts = Counter(hashlib.sha1(blob).hexdigest() for _, blob, _, _ in items)
    found, seen = [], set()
    for page, blob, context, caption in items:
        digest = hashlib.sha1(blob).hexdigest()
        if digest in seen or counts[digest] >= REPEATED_ON:
            continue                                       # duplicate, or a logo on every slide
        seen.add(digest)
        try:
            import io
            with Image.open(io.BytesIO(blob)) as im:
                width, height = im.size
            if min(width, height) < MIN_IMAGE_PX:
                continue
            png = _to_png(blob)
        except Exception:
            continue                                       # EMF/WMF or broken image: skip
        found.append({"page": page, "png": png, "width": width, "height": height,
                      "context": context, "caption": caption})
        if len(found) >= max_images:
            break
    return found


def location(source: str, page) -> str:
    """', slide 3' for PowerPoint, ', page 3' otherwise, '' when there's no page."""
    if not page:
        return ""
    return f", {'slide' if str(source).lower().endswith('.pptx') else 'page'} {page}"
